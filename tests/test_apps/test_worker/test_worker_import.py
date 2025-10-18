from __future__ import annotations

import io
from http import HTTPStatus

import openpyxl
import pytest

from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.exceptions import ValidationError


from server.apps.workers.models import Worker
from server.apps.workers.serializers import WorkerImportSerializer
from server.apps.workers.services.import_workers import WorkerImportService


@pytest.mark.django_db
def test_import_from_excel_success() -> None:
    """Тест успешного импорта работников из Excel."""

    # Создаём временный Excel-файл в памяти
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['first_name', 'last_name', 'email', 'position'])
    sheet.append(['John', 'Doe', 'john@example.com', 'guard'])
    sheet.append(['Jane', 'Smith', 'jane@example.com', 'other'])

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    # Django UploadedFile-обёртка
    file = SimpleUploadedFile(
        'test_workers.xlsx',
        stream.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )

    service = WorkerImportService()
    result = service.import_from_excel(file)

    assert result['created'] == 2
    assert result['errors'] == []
    assert Worker.objects.count() == 2


@pytest.mark.django_db
def test_import_from_excel_with_errors() -> None:
    """Тест обработки ошибок при импорте."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['first_name', 'last_name', 'email', 'position'])
    sheet.append(['', 'Doe', 'noemail', 'other'])
    sheet.append(['John', 'Doe', 'bad_email', 'other'])

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    file = SimpleUploadedFile(
        'test_workers.xlsx',
        stream.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )

    service = WorkerImportService()
    result = service.import_from_excel(file)

    assert result['created'] == 0
    assert len(result['errors']) == 2
    assert 'Required fields are missin' in str(result['errors'][0]['error']) or \
           'Incorrect email' in str(result['errors'][1]['error'])


@pytest.mark.django_db
def test_import_from_excel_duplicate_email(worker_factory) -> None:
    """Тест, что дубликат email не создаётся повторно."""
    worker_factory(email='exists@test.com')

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['first_name', 'last_name', 'email', 'position'])
    sheet.append(['John', 'Doe', 'exists@test.com', 'other'])

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    file = SimpleUploadedFile(
        'test_workers.xlsx',
        stream.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )

    service = WorkerImportService()
    result = service.import_from_excel(file)

    assert result['created'] == 0
    assert len(result['errors']) == 1
    assert 'Email already exists' in result['errors'][0]['error']


@pytest.mark.django_db
def test_import_from_excel_invalid_file() -> None:
    """Тест: обработка невалидного Excel файла."""
    # создаем "битый" файл — не настоящий Excel
    invalid_file = SimpleUploadedFile(
        'not_excel.txt',
        b'not an excel content',
        content_type='text/plain',
    )

    service = WorkerImportService()
    result = service.import_from_excel(invalid_file)

    assert result == {
        'created': 0,
        'errors': [{'error': 'Excel file cannot be read.'}],
    }


def test_validate_xl_file_valid() -> None:
    """Проверяет, что .xlsx файл проходит валидацию."""
    excel_file = SimpleUploadedFile("test.xlsx", b"dummy content")
    serializer = WorkerImportSerializer(data={"xl_file": excel_file})
    assert serializer.is_valid(), serializer.errors


def test_validate_xl_file_invalid_extension() -> None:
    """Проверяет, что не .xlsx файл вызывает ValidationError."""
    not_excel_file = SimpleUploadedFile("test.txt", b"dummy content")
    serializer = WorkerImportSerializer()

    with pytest.raises(
        ValidationError, match="Файл должен быть в формате .xlsx"
    ):
        serializer.validate_xl_file(not_excel_file)


@pytest.mark.django_db
def test_import_view_no_file(auth_admin_client) -> None:
    """Проверяет, что без файла возвращается 400."""
    response = auth_admin_client.post('/api/workers/import/', data={})

    assert response.status_code == HTTPStatus.BAD_REQUEST
    assert response.data == {'error': 'File was not transferred.'}


@pytest.mark.django_db
def test_import_view_with_valid_file(auth_admin_client) -> None:
    """Проверяет успешную загрузку Excel файла."""
    # Создаём временный Excel-файл в памяти
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(['first_name', 'last_name', 'email', 'position'])
    sheet.append(['John', 'Doe', 'john@example.com', 'cleaner'])

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    file = SimpleUploadedFile(
        "test.xlsx",
        stream.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response = auth_admin_client.post(
        '/api/workers/import/',
        data={'file': file},
        format='multipart',
    )

    assert response.status_code == HTTPStatus.OK
    assert response.data['created'] == 1
    assert response.data['errors'] == []
    assert Worker.objects.count() == 1